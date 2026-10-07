"""Læs Excel og byg organisationshierarkier i memory."""

import logging
from collections import defaultdict, deque
from copy import deepcopy
from datetime import datetime
from io import BytesIO
from urllib.parse import urljoin, urlsplit
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
import configuration as config

LOGGER = logging.getLogger(__name__)
_CATALOGUE = None


def _read_roots(content, year):
    """Returnerer fanenavn og en dictionary med navn -> type."""
    sheet = (
        config.ORGANISATIONS_SHEET_2026
        if year == 2026
        else config.ORGANISATIONS_SHEET_OTHER
    )

    workbook = load_workbook(
        BytesIO(content),
        read_only=True,
        data_only=True,
    )

    try:
        if sheet not in workbook.sheetnames:
            raise ValueError(f"Excel-fanen {sheet!r} mangler.")

        rows = workbook[sheet].iter_rows(
            max_col=2,
            values_only=True,
        )

        if next(rows, None) != (
            "Overordnet organisationsnavn",
            "Type",
        ):
            raise ValueError(
                "Forventede organisationsnavn i A1 og Type i B1."
            )

        roots = {}

        for row_number, (name, category) in enumerate(rows, start=2):
            if name is None and category is None:
                continue

            if not all(
                isinstance(value, str) and value.strip()
                for value in (name, category)
            ):
                raise ValueError(
                    f"Excel-række {row_number} mangler navn eller type."
                )

            name = name.strip()
            category = category.strip()

            if name in roots and roots[name] != category:
                raise ValueError(
                    f"Modstridende typer for {name!r}."
                )

            roots[name] = category

        if not roots:
            raise ValueError(
                "Excel-fanen indeholder ingen organisationer."
            )

        return sheet, roots

    finally:
        workbook.close()


def _fetch_organizations():
    """Henter rå organisationer og følger annoncerede næste sider."""
    from q_cura_api import api_client
    from q_cura_api.functionality.organisation_hent import (
        get_organizations,
    )

    bundle = get_organizations(
        search_name=None,
        raw=True,
        include_inactive=True,
    )

    organizations = {}
    seen_urls = set()

    base = api_client.BASE_URL
    base_parts = urlsplit(base)

    for page_number in range(
        1,
        config.ORGANISATIONS_MAX_PAGES + 1,
    ):
        if (
            not isinstance(bundle, dict)
            or bundle.get("resourceType") != "Bundle"
        ):
            raise RuntimeError(
                "Organisationsopslaget returnerede ikke et Bundle."
            )

        entries = bundle.get("entry", [])
        links = bundle.get("link", [])

        if not isinstance(entries, list) or not isinstance(links, list):
            raise RuntimeError("Uventet entry/link-format.")

        for entry in entries:
            if (
                not isinstance(entry, dict)
                or not isinstance(entry.get("resource"), dict)
            ):
                raise RuntimeError("Uventet organisationsresource.")

            resource = entry["resource"]

            if resource.get("resourceType") != "Organization":
                continue

            org_id = resource.get("id")

            if not isinstance(org_id, str) or not org_id.strip():
                raise RuntimeError("En organisation mangler id.")

            if (
                org_id in organizations
                and organizations[org_id] != resource
            ):
                raise RuntimeError(
                    "Samme organisations-id har forskellige data."
                )

            organizations[org_id] = resource

        if any(not isinstance(link, dict) for link in links):
            raise RuntimeError("Uventet pagination-link.")

        next_links = [
            link.get("url")
            for link in links
            if link.get("relation") == "next"
        ]

        LOGGER.info("Organisationsside %s læst", page_number)

        if not next_links:
            return organizations

        if (
            len(next_links) != 1
            or not isinstance(next_links[0], str)
            or not next_links[0]
        ):
            raise RuntimeError("Ugyldigt next-link.")

        url = urljoin(base, next_links[0])
        parts = urlsplit(url)

        # Send aldrig CURA-kald til et andet miljø eller en anden vært.
        if (
            parts.scheme != base_parts.scheme
            or parts.netloc != base_parts.netloc
            or not url.startswith(base)
            or parts.fragment
            or url in seen_urls
        ):
            raise RuntimeError(
                "Ugyldigt eller gentaget pagination-link."
            )

        if page_number == config.ORGANISATIONS_MAX_PAGES:
            raise RuntimeError(
                "Sidegrænsen er nået; listen er ikke færdig."
            )

        seen_urls.add(url)

        bundle = api_client.get(
            url[len(base):],
            raw=True,
        )

    raise RuntimeError(
        "Ingen afslutning på organisationsopslaget."
    )


def _build_catalogue(roots, organizations):
    """Finder rødder og alle efterkommere. Returnerer en unik liste."""
    children = defaultdict(list)

    # Byg et indeks: forældrereference -> børnenes id'er.
    for org_id, org in organizations.items():
        part_of = org.get("partOf") or {}

        if not isinstance(part_of, dict):
            raise RuntimeError("Uventet partOf-format.")

        reference = part_of.get("reference")

        if reference:
            if (
                not isinstance(reference, str)
                or not reference.startswith("Organization/")
            ):
                raise RuntimeError("Ukendt partOf-reference.")

            children[reference].append(org_id)

    result = {}

    for root_name, category in roots.items():
        matches = [
            org_id
            for org_id, org in organizations.items()
            if org.get("name") == root_name
        ]

        if len(matches) != 1:
            raise ValueError(
                f"Forventede ét præcist match på {root_name!r}, "
                f"fandt {len(matches)}."
            )

        pending = deque([matches[0]])
        visited = set()

        # Fortsæt indtil alle niveauer under denne rod er gennemgået.
        while pending:
            org_id = pending.popleft()

            if org_id in visited:
                raise RuntimeError(
                    "Cirkulært organisationshierarki fundet."
                )

            visited.add(org_id)
            org = organizations[org_id]

            pending.extend(
                children.get("Organization/" + org_id, [])
            )

            # Gennemgå også inaktive forældre, så deres børn ikke mistes.
            if (
                not config.ORGANISATIONS_INCLUDE_INACTIVE
                and org.get("active") is not True
            ):
                continue

            if org_id in result:
                if (
                    result[org_id]["type"].casefold()
                    != category.casefold()
                ):
                    raise ValueError(
                        "En organisation arver modstridende typer."
                    )

                result[org_id]["root_names"].append(root_name)
                continue

            result[org_id] = {
                "organization_id": org_id,
                "name": org.get("name"),
                "active": org.get("active"),
                "part_of_reference": (
                    org.get("partOf") or {}
                ).get("reference"),
                "type": category,
                "root_names": [root_name],
            }

    return sorted(
        result.values(),
        key=lambda org: (
            str(org["name"] or "").casefold(),
            org["organization_id"],
        ),
    )


def initialize_organisations():
    """Genindlæser kataloget før første item. Fejl stopper opstarten."""
    global _CATALOGUE

    # Et mislykket opslag må ikke efterlade et gammelt katalog.
    _CATALOGUE = None

    if type(config.ORGANISATIONS_INCLUDE_INACTIVE) is not bool:
        raise ValueError(
            "ORGANISATIONS_INCLUDE_INACTIVE skal være bool."
        )

    if (
        type(config.ORGANISATIONS_MAX_PAGES) is not int
        or config.ORGANISATIONS_MAX_PAGES < 1
    ):
        raise ValueError(
            "ORGANISATIONS_MAX_PAGES skal være et positivt heltal."
        )

    from q_sharepoint_api.sp_api import get_client
    from q_cura_api.api_client import set_cura_credential

    client = get_client()
    site_id = client.get_site_id(config.ORGANISATIONS_SITE)

    downloaded = client.download_file_to_memory_by_path(
        site_id=site_id,
        file_path=config.ORGANISATIONS_FILE_PATH,
        save_dir=None,
    )

    year = datetime.now(
        ZoneInfo(config.SEARCH_TIMEZONE)
    ).year

    sheet, roots = _read_roots(
        downloaded["file_bytes"],
        year,
    )

    set_cura_credential(config.CURA_CREDENTIAL_NAME)
    organizations = _fetch_organizations()

    _CATALOGUE = {
        "sheet": sheet,
        "loaded_at": datetime.now(
            ZoneInfo(config.SEARCH_TIMEZONE)
        ).isoformat(),
        "organizations": _build_catalogue(
            roots,
            organizations,
        ),
    }

    LOGGER.info(
        "Katalog klar: fane=%s, organisationer=%s",
        sheet,
        len(_CATALOGUE["organizations"]),
    )

    return deepcopy(_CATALOGUE)


def get_organisations():
    """Returnerer en kopi af kataloget uden nye API-kald."""
    if _CATALOGUE is None:
        raise RuntimeError(
            "Organisationskataloget er ikke indlæst."
        )

    return deepcopy(_CATALOGUE)